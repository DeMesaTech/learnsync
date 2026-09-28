import tempfile
import unittest
from pathlib import Path

from docx import Document

from module_extractor import extract_syllabus_topics


def set_cell_paragraphs(cell, paragraphs):
    cell.text = ""
    for index, (text, bold) in enumerate(paragraphs):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.add_run(text).bold = bold


class SyllabusChapterExtractionTests(unittest.TestCase):
    def test_extracts_chapters_from_topics_column_not_ilo_column(self):
        document = Document()
        main_table = document.add_table(rows=5, cols=3)
        for row in main_table.rows:
            row.cells[0].text = "Syllabus preamble"

        header = main_table.add_row().cells
        header[0].text = "Time Allotment"
        header[1].text = "Intended Learning Outcomes (ILO)"
        header[2].text = "Topics"

        week_one = main_table.add_row().cells
        week_one[0].text = "Week 1"
        week_one[1].text = "Students will understand the course goals."
        set_cell_paragraphs(week_one[2], [("COURSE ORIENTATION", True)])

        week_two = main_table.add_row().cells
        week_two[0].text = "Week 2"
        week_two[1].text = "Identify and explain current and voltage."
        set_cell_paragraphs(week_two[2], [
            ("CHAPTER 1 - Fundamentals of Current and Voltage", True),
            ("Electrical Quantities", False),
        ])

        week_three = main_table.add_row().cells
        week_three[0].text = "Week 3"
        week_three[1].text = "Understand power supply units."
        set_cell_paragraphs(week_three[2], [
            ("CHAPTER 2 - Basics of Power", True),
            ("Supply Unit", True),
            ("Transformer Circuit", False),
        ])

        continuation_table = document.add_table(rows=0, cols=4)
        week_four = continuation_table.add_row().cells
        week_four[0].text = "Week 4"
        week_four[2].text = "Identify semiconductor types."
        set_cell_paragraphs(week_four[3], [
            ("CHAPTER 3 - Semiconductors", True),
            ("Types of Materials", False),
        ])

        with tempfile.TemporaryDirectory() as directory:
            syllabus_path = Path(directory) / "syllabus.docx"
            document.save(syllabus_path)

            topics = extract_syllabus_topics(syllabus_path)

        self.assertEqual(topics, [
            "CHAPTER 1 - Fundamentals of Current and Voltage",
            "CHAPTER 2 - Basics of Power Supply Unit",
            "CHAPTER 3 - Semiconductors",
        ])


if __name__ == "__main__":
    unittest.main()
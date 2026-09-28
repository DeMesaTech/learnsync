"""Faculty reviewed topic content and student outline."""
import html
import re
import uuid
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Literal

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field
from psycopg2.extras import RealDictCursor
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

from db import get_db_connection

content_router = APIRouter(prefix="/api/learning", tags=["learning content"])


class SafeHTML(HTMLParser):
    allowed = {"p", "br", "strong", "em", "b", "i", "h2", "h3", "h4", "ul", "ol", "li", "a"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag not in self.allowed:
            return
        if tag == "a":
            href = next((value for key, value in attrs if key == "href"), "")
            if href.startswith(("https://", "http://")):
                self.parts.append(f'<a href="{html.escape(href, quote=True)}" rel="noopener noreferrer">')
            else:
                self.parts.append("<a>")
        else:
            self.parts.append(f"<{tag}>")

    def handle_endtag(self, tag):
        if tag in self.allowed and tag != "br":
            self.parts.append(f"</{tag}>")

    def handle_data(self, data):
        self.parts.append(html.escape(data))


def sanitize(markup):
    parser = SafeHTML()
    parser.feed(markup or "")
    return "".join(parser.parts)


def has_visible_text(markup):
    return bool(html.unescape(re.sub(r"<[^>]*>", " ", markup or "")).replace("\u200b", "").strip())


class ContentInput(BaseModel):
    teacher_id: int
    topic_key: str = Field(min_length=1, max_length=64)
    title: str = Field(min_length=1, max_length=255)
    edited_html: str = ""
    publish: bool = False


class ResourceInput(BaseModel):
    teacher_id: int
    kind: Literal["module", "reference", "material"]
    content_level: Literal["chapter", "topic"]
    content_key: str = Field(min_length=1, max_length=64)
    title: str = Field(min_length=1, max_length=255)
    body_html: str = ""
    url: str | None = None
    publish: bool = False


class ActivityInput(BaseModel):
    teacher_id: int
    title: str = Field(min_length=1, max_length=255)
    description: str = ""
    due_date: datetime | None = None
    points: float = Field(gt=0)
    grading_period: Literal["Midterm", "Finals"] = "Midterm"
    delivery_type: Literal["online", "offline"] = "online"
    content_level: Literal["course", "chapter", "subsection", "topic"] = "course"
    content_key: str | None = None
    section_ids: list[int] = Field(min_length=1)
    publish: bool = False


class OfflineActivityScore(BaseModel):
    student_id: int
    score: float | None = Field(default=None, ge=0)


class OfflineActivityScores(BaseModel):
    teacher_id: int
    scores: list[OfflineActivityScore]


def _teacher(cur, class_id, teacher_id):
    cur.execute("SELECT 1 FROM class WHERE class_id=%s AND employee_id=%s AND workflow_status='active'",
                (class_id, teacher_id))
    if not cur.fetchone():
        raise HTTPException(status_code=403, detail="Assigned faculty member required.")


def _outline(cur, class_id):
    cur.execute("SELECT content,version FROM syllabus_version WHERE class_id=%s AND status='approved' ORDER BY version DESC LIMIT 1", (class_id,))
    row = cur.fetchone()
    if not row:
        raise HTTPException(status_code=409, detail="Approve a syllabus first.")
    return row


def _topic(cur, class_id, topic_key):
    outline = _outline(cur, class_id)["content"]
    for chapter in outline.get("chapters", []):
        topics = list(chapter.get("topics") or [])
        for subsection in chapter.get("subsections") or []:
            topics.extend(subsection.get("topics") or [])
        if any(topic.get("key") == topic_key for topic in topics):
            return
    raise HTTPException(status_code=422, detail="Choose a topic in the approved syllabus.")


def _location(cur, class_id, level, key):
    content = _outline(cur, class_id)["content"]
    for chapter in content.get("chapters", []):
        if level == "chapter" and chapter.get("key") == key:
            return
        topics = list(chapter.get("topics") or [])
        for subsection in chapter.get("subsections") or []:
            topics.extend(subsection.get("topics") or [])
        if level == "topic" and any(topic.get("key") == key for topic in topics):
            return
    raise HTTPException(status_code=422, detail="Choose a chapter or topic in the approved syllabus.")


@content_router.get("/{class_id}/outline")
def outline(class_id: int, student_id: int | None = None, teacher_id: int | None = None):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        if teacher_id is not None:
            _teacher(cur, class_id, teacher_id)
        elif student_id is not None:
            cur.execute("""SELECT 1 FROM enrollment e JOIN class c ON c.class_id=e.class_id
                           WHERE e.class_id=%s AND e.student_id=%s AND c.workflow_status='active'""", (class_id, student_id))
            if not cur.fetchone():
                raise HTTPException(status_code=403, detail="Subject enrollment required.")
        else:
            raise HTTPException(status_code=403, detail="Sign in to view this outline.")
        try:
            syllabus = _outline(cur, class_id)
        except HTTPException as error:
            if teacher_id is None and error.status_code == 409:
                return {"status": "pending", "message": "Syllabus pending faculty approval.",
                        "content": {"chapters": []}, "materials": [], "resources": [],
                        "activities": [], "quizzes": []}
            raise
        cur.execute("""SELECT content_id,topic_key,title,source_type,file_path,extracted_text,edited_html,status,published_at
                       FROM learning_content WHERE class_id=%s AND (%s IS NOT NULL OR status='published')
                       ORDER BY created_at,content_id""", (class_id, teacher_id))
        contents = cur.fetchall()
        if teacher_id is None:
            for item in contents:
                item.pop("extracted_text", None)
        cur.execute("""SELECT q.quiz_id,q.title,q.date_created,q.delivery_type,q.status,q.content_level,q.content_key,
                              q.total_points,q.grading_period,q.origin,
                              ARRAY(SELECT qs.section_id FROM quiz_sections qs WHERE qs.quiz_id=q.quiz_id) AS section_ids
                       FROM quiz q WHERE q.class_id=%s AND (%s IS NOT NULL OR q.status='Published')
                       AND (%s IS NOT NULL OR NOT EXISTS (SELECT 1 FROM quiz_sections qs WHERE qs.quiz_id=q.quiz_id)
                         OR EXISTS (SELECT 1 FROM enrollment e JOIN quiz_sections qs ON qs.section_id=e.section_id
                         WHERE e.student_id=%s AND e.class_id=q.class_id AND qs.quiz_id=q.quiz_id))
                       ORDER BY q.quiz_id""", (class_id, teacher_id, teacher_id, student_id))
        quizzes = cur.fetchall()
        cur.execute("""SELECT resource_id,kind,content_level,content_key,title,body_html,url,
                              file_path,original_filename,extracted_text,status,display_order,legacy_module_id
                       FROM learning_resource WHERE class_id=%s AND (%s IS NOT NULL OR status='published')
                       ORDER BY display_order,resource_id""", (class_id, teacher_id))
        resources = cur.fetchall()
        if teacher_id is None:
            for resource in resources:
                resource.pop("extracted_text", None)
        cur.execute("""SELECT a.activity_id,a.title,a.description,a.due_date,a.points,a.status,
                              a.delivery_type,a.content_level,a.content_key,a.grading_period,
                              ARRAY(SELECT x.section_id FROM activity_sections x WHERE x.activity_id=a.activity_id) AS section_ids
                       FROM activity a WHERE a.class_id=%s AND (%s IS NOT NULL OR a.status='Published')
                         AND (%s IS NOT NULL OR NOT EXISTS
                           (SELECT 1 FROM activity_sections x WHERE x.activity_id=a.activity_id)
                           OR EXISTS (SELECT 1 FROM activity_sections x JOIN enrollment e
                             ON e.section_id=x.section_id AND e.class_id=a.class_id
                             WHERE x.activity_id=a.activity_id AND e.student_id=%s))
                       ORDER BY a.activity_id""", (class_id, teacher_id, teacher_id, student_id))
        activities = cur.fetchall()
        return {"status": "approved", "version": syllabus["version"], "content": syllabus["content"],
                "materials": contents, "resources": resources, "activities": activities, "quizzes": quizzes}
    finally:
        conn.close()


@content_router.post("/{class_id}/contents")
def save_content(class_id: int, payload: ContentInput):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, payload.teacher_id)
        _topic(cur, class_id, payload.topic_key)
        markup = sanitize(payload.edited_html)
        if payload.publish and not has_visible_text(markup):
            raise HTTPException(status_code=422, detail="Review and enter content before publishing.")
        cur.execute("""INSERT INTO learning_content(class_id,topic_key,title,source_type,edited_html,status,created_by,published_at)
                       VALUES (%s,%s,%s,'manual',%s,%s,%s,CASE WHEN %s THEN CURRENT_TIMESTAMP END)
                       RETURNING content_id""", (class_id, payload.topic_key, payload.title.strip(), markup,
                                                 "published" if payload.publish else "draft", payload.teacher_id,
                                                 payload.publish))
        result = cur.fetchone()
        conn.commit()
        return result
    finally:
        conn.close()


@content_router.post("/{class_id}/contents/pdf")
async def upload_pdf(class_id: int, teacher_id: int = Form(...), topic_key: str = Form(...),
                     title: str = Form(...), file: UploadFile = File(...)):
    if Path(file.filename or "").suffix.lower() != ".pdf":
        raise HTTPException(status_code=415, detail="Upload a PDF.")
    data = await file.read(20 * 1024 * 1024 + 1)
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="PDF must be 20 MB or smaller.")
    conn = get_db_connection()
    path = None
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, teacher_id)
        _topic(cur, class_id, topic_key)
        folder = Path(__file__).resolve().parents[1] / "uploads" / "content"
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{uuid.uuid4().hex}.pdf"
        path.write_bytes(data)
        reader = PdfReader(path)
        extracted = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        cur.execute("""INSERT INTO learning_content(class_id,topic_key,title,source_type,file_path,extracted_text,status,created_by)
                       VALUES (%s,%s,%s,'pdf',%s,%s,'draft',%s) RETURNING content_id""",
                    (class_id, topic_key, title.strip(), f"uploads/content/{path.name}", extracted, teacher_id))
        result = cur.fetchone()
        conn.commit()
        return {**result, "extracted_text": extracted, "needs_review": True}
    except Exception:
        conn.rollback()
        if path:
            path.unlink(missing_ok=True)
        raise
    finally:
        conn.close()


@content_router.put("/{class_id}/contents/{content_id}")
def review_content(class_id: int, content_id: int, payload: ContentInput):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, payload.teacher_id)
        _topic(cur, class_id, payload.topic_key)
        markup = sanitize(payload.edited_html)
        if payload.publish and not has_visible_text(markup):
            raise HTTPException(status_code=422, detail="Review and enter content before publishing.")
        cur.execute("""UPDATE learning_content SET topic_key=%s,title=%s,edited_html=%s,status=%s,
                       published_at=CASE WHEN %s THEN CURRENT_TIMESTAMP ELSE NULL END
                       WHERE class_id=%s AND content_id=%s RETURNING content_id""",
                    (payload.topic_key, payload.title.strip(), markup,
                     "published" if payload.publish else "draft", payload.publish, class_id, content_id))
        result = cur.fetchone()
        if not result:
            raise HTTPException(status_code=404, detail="Content not found.")
        conn.commit()
        return result
    finally:
        conn.close()


@content_router.post("/{class_id}/resources", status_code=201)
def create_resource(class_id: int, payload: ResourceInput):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, payload.teacher_id)
        _location(cur, class_id, payload.content_level, payload.content_key)
        if payload.kind == "module" and payload.content_level != "chapter":
            raise HTTPException(status_code=422, detail="Lesson modules belong to a chapter.")
        if payload.kind == "material":
            raise HTTPException(status_code=422, detail="Upload a file to create a learning material.")
        if payload.kind == "reference" and payload.url and not payload.url.startswith(("https://", "http://")):
            raise HTTPException(status_code=422, detail="Reference links must use HTTP or HTTPS.")
        markup = sanitize(payload.body_html)
        if payload.publish and payload.kind == "module" and not has_visible_text(markup):
            raise HTTPException(status_code=422, detail="Add lesson text before publishing.")
        cur.execute("""INSERT INTO learning_resource(class_id,kind,content_level,content_key,title,
                       body_html,url,status,created_by,published_at)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,CASE WHEN %s THEN CURRENT_TIMESTAMP END)
                       RETURNING resource_id""",
                    (class_id, payload.kind, payload.content_level, payload.content_key,
                     payload.title.strip(), markup, payload.url,
                     "published" if payload.publish else "draft", payload.teacher_id, payload.publish))
        result = cur.fetchone()
        conn.commit()
        return result
    finally:
        conn.close()


@content_router.put("/{class_id}/resources/{resource_id}")
def update_resource(class_id: int, resource_id: int, payload: ResourceInput):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, payload.teacher_id)
        _location(cur, class_id, payload.content_level, payload.content_key)
        cur.execute("SELECT kind,file_path FROM learning_resource WHERE class_id=%s AND resource_id=%s",
                    (class_id, resource_id))
        existing = cur.fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="Resource not found.")
        if existing["kind"] != payload.kind:
            raise HTTPException(status_code=422, detail="Resource type cannot be changed.")
        if payload.kind == "module" and payload.content_level != "chapter":
            raise HTTPException(status_code=422, detail="Lesson modules belong to a chapter.")
        if payload.kind == "reference" and payload.url and not payload.url.startswith(("https://", "http://")):
            raise HTTPException(status_code=422, detail="Reference links must use HTTP or HTTPS.")
        markup = sanitize(payload.body_html)
        if payload.publish and payload.kind == "module" and not has_visible_text(markup) and not existing["file_path"]:
            raise HTTPException(status_code=422, detail="Add lesson text or a file before publishing.")
        cur.execute("""UPDATE learning_resource SET content_level=%s,content_key=%s,title=%s,
                       body_html=%s,url=%s,status=%s,published_at=CASE WHEN %s THEN CURRENT_TIMESTAMP ELSE NULL END
                       WHERE resource_id=%s RETURNING resource_id""",
                    (payload.content_level, payload.content_key, payload.title.strip(), markup,
                     payload.url, "published" if payload.publish else "draft", payload.publish, resource_id))
        result = cur.fetchone()
        conn.commit()
        return result
    finally:
        conn.close()


@content_router.delete("/{class_id}/resources/{resource_id}")
def delete_resource(class_id: int, resource_id: int, teacher_id: int):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, teacher_id)
        cur.execute("DELETE FROM learning_resource WHERE class_id=%s AND resource_id=%s RETURNING file_path,legacy_module_id",
                    (class_id, resource_id))
        removed = cur.fetchone()
        if not removed:
            raise HTTPException(status_code=404, detail="Learning material or reference not found.")
        stored_file = removed["file_path"]
        may_remove = bool(stored_file and not removed["legacy_module_id"])
        if may_remove:
            cur.execute("SELECT 1 FROM learning_resource WHERE file_path=%s LIMIT 1", (stored_file,))
            may_remove = not bool(cur.fetchone())
        conn.commit()
        if may_remove:
            upload_dir = (Path(__file__).resolve().parents[1] / "uploads" / "learning").resolve()
            target = (Path(__file__).resolve().parents[1] / stored_file).resolve()
            if target.is_relative_to(upload_dir):
                try:
                    target.unlink(missing_ok=True)
                except OSError:
                    pass
        return {"deleted": True}
    finally:
        conn.close()


@content_router.delete("/{class_id}/contents/{content_id}")
def delete_content(class_id: int, content_id: int, teacher_id: int):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, teacher_id)
        cur.execute("DELETE FROM learning_content WHERE class_id=%s AND content_id=%s RETURNING content_id",
                    (class_id, content_id))
        if not cur.fetchone():
            raise HTTPException(status_code=404, detail="Topic content not found.")
        conn.commit()
        return {"deleted": True}
    finally:
        conn.close()


@content_router.post("/{class_id}/resources/upload", status_code=201)
async def upload_resource(class_id: int, teacher_id: int = Form(...),
                          kind: Literal["module", "material"] = Form(...),
                          content_level: Literal["chapter", "topic"] = Form(...),
                          content_key: str = Form(...), title: str = Form(...),
                          file: UploadFile = File(...)):
    allowed = {".pdf", ".docx", ".pptx", ".png", ".jpg", ".jpeg"}
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in allowed:
        raise HTTPException(status_code=415, detail="Upload PDF, DOCX, PPTX, PNG, or JPEG.")
    data = await file.read(20 * 1024 * 1024 + 1)
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File must be 20 MB or smaller.")
    if kind == "module" and content_level != "chapter":
        raise HTTPException(status_code=422, detail="Lesson modules belong to a chapter.")
    conn = get_db_connection()
    path = None
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, teacher_id)
        _location(cur, class_id, content_level, content_key)
        folder = Path(__file__).resolve().parents[1] / "uploads" / "learning"
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{uuid.uuid4().hex}{suffix}"
        path.write_bytes(data)
        extracted = ""
        try:
            if suffix == ".pdf":
                extracted = "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
            elif suffix == ".docx":
                extracted = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
            elif suffix == ".pptx":
                presentation = Presentation(path)
                extracted = "\n".join(shape.text for slide in presentation.slides
                                      for shape in slide.shapes if shape.has_text_frame)
        except Exception as error:
            raise HTTPException(status_code=422, detail="The uploaded document could not be read.") from error
        cur.execute("""INSERT INTO learning_resource(class_id,kind,content_level,content_key,title,
                       file_path,original_filename,extracted_text,status,created_by)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s,'draft',%s) RETURNING resource_id""",
                    (class_id, kind, content_level, content_key, title.strip(),
                     f"uploads/learning/{path.name}", Path(file.filename or "file").name,
                     extracted, teacher_id))
        result = cur.fetchone()
        conn.commit()
        return result
    except Exception:
        conn.rollback()
        if path:
            path.unlink(missing_ok=True)
        raise
    finally:
        conn.close()


@content_router.post("/{class_id}/resources/import-legacy")
def import_legacy_modules(class_id: int, teacher_id: int):
    """Match historical module uploads only when a chapter title is unique."""
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, teacher_id)
        chapters = _outline(cur, class_id)["content"].get("chapters", [])
        by_title = {}
        for chapter in chapters:
            by_title.setdefault(chapter.get("title", "").strip().casefold(), []).append(chapter)
        cur.execute("""SELECT m.module_id,m.title,m.summary,m.file_path,t.title AS chapter_title
                       FROM module m LEFT JOIN syllabus_topic t ON t.topic_id=m.topic_id
                       WHERE m.class_id=%s AND NOT EXISTS
                       (SELECT 1 FROM learning_resource r WHERE r.legacy_module_id=m.module_id)
                       ORDER BY m.module_id""", (class_id,))
        unmatched = []
        imported = 0
        for module in cur.fetchall():
            matches = by_title.get((module["chapter_title"] or "").strip().casefold(), [])
            if len(matches) != 1:
                unmatched.append(dict(module))
                continue
            cur.execute("""INSERT INTO learning_resource(class_id,kind,content_level,content_key,title,
                           body_html,file_path,original_filename,status,created_by,legacy_module_id)
                           VALUES (%s,'module','chapter',%s,%s,%s,%s,%s,'draft',%s,%s)""",
                        (class_id, matches[0]["key"], module["title"] or "Untitled lesson",
                         sanitize(f"<p>{html.escape(module['summary'] or '')}</p>"), module["file_path"],
                         Path(module["file_path"] or "file").name, teacher_id, module["module_id"]))
            imported += 1
        conn.commit()
        return {"imported": imported, "unmatched": unmatched}
    finally:
        conn.close()


@content_router.post("/{class_id}/activities", status_code=201)
def create_activity(class_id: int, payload: ActivityInput):
    from routers.faculty_work import _valid_content_placement, _section
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, payload.teacher_id)
        _valid_content_placement(cur, class_id, payload.content_level, payload.content_key)
        for section_id in set(payload.section_ids):
            _section(cur, class_id, section_id)
        cur.execute("""INSERT INTO activity(class_id,title,description,due_date,employee_id,points,
                       status,grading_period,delivery_type,content_level,content_key)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING activity_id""",
                    (class_id, payload.title.strip(), payload.description, payload.due_date,
                     payload.teacher_id, payload.points, "Published" if payload.publish else "Draft",
                     payload.grading_period, payload.delivery_type, payload.content_level, payload.content_key))
        activity_id = cur.fetchone()["activity_id"]
        for section_id in set(payload.section_ids):
            cur.execute("INSERT INTO activity_sections(activity_id,section_id) VALUES (%s,%s)",
                        (activity_id, section_id))
            if payload.publish:
                cur.execute("DELETE FROM grade_publication WHERE class_id=%s AND section_id=%s",
                            (class_id, section_id))
        conn.commit()
        return {"activity_id": activity_id, "status": "Published" if payload.publish else "Draft"}
    finally:
        conn.close()


@content_router.delete("/{class_id}/activities/{activity_id}")
def delete_activity(class_id: int, activity_id: int, teacher_id: int):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, teacher_id)
        cur.execute("SELECT status FROM activity WHERE class_id=%s AND activity_id=%s", (class_id, activity_id))
        activity = cur.fetchone()
        if not activity:
            raise HTTPException(status_code=404, detail="Activity not found.")
        cur.execute("SELECT 1 FROM act_submission WHERE activity_id=%s LIMIT 1", (activity_id,))
        if cur.fetchone():
            raise HTTPException(status_code=409, detail="This activity has submissions and must be retained for grading history.")
        cur.execute("DELETE FROM activity_sections WHERE activity_id=%s", (activity_id,))
        cur.execute("DELETE FROM activity WHERE class_id=%s AND activity_id=%s RETURNING activity_id",
                    (class_id, activity_id))
        if not cur.fetchone():
            raise HTTPException(status_code=404, detail="Activity not found.")
        cur.execute("DELETE FROM syllabus_item_link WHERE class_id=%s AND item_type='activity' AND item_id=%s",
                    (class_id, activity_id))
        if activity["status"] == "Published":
            cur.execute("DELETE FROM grade_publication WHERE class_id=%s", (class_id,))
        conn.commit()
        return {"deleted": True}
    finally:
        conn.close()


@content_router.patch("/{class_id}/activities/{activity_id}/status")
def set_activity_status(class_id: int, activity_id: int, teacher_id: int,
                        status: Literal["Draft", "Published"]):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, teacher_id)
        cur.execute("""UPDATE activity SET status=%s,publish_date=CASE WHEN %s='Published'
                       THEN CURRENT_TIMESTAMP ELSE NULL END
                       WHERE class_id=%s AND activity_id=%s RETURNING activity_id""",
                    (status, status, class_id, activity_id))
        if not cur.fetchone():
            raise HTTPException(status_code=404, detail="Activity not found.")
        cur.execute("""DELETE FROM grade_publication WHERE class_id=%s AND section_id IN
                       (SELECT section_id FROM activity_sections WHERE activity_id=%s)""",
                    (class_id, activity_id))
        conn.commit()
        return {"activity_id": activity_id, "status": status}
    finally:
        conn.close()


@content_router.put("/{class_id}/activities/{activity_id}/offline-scores")
def save_offline_activity_scores(class_id: int, activity_id: int, payload: OfflineActivityScores):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _teacher(cur, class_id, payload.teacher_id)
        cur.execute("""SELECT points,grading_period FROM activity WHERE activity_id=%s
                       AND class_id=%s AND delivery_type='offline'""", (activity_id, class_id))
        activity = cur.fetchone()
        if not activity:
            raise HTTPException(status_code=404, detail="Offline activity not found.")
        cur.execute("""SELECT e.student_id,e.section_id FROM enrollment e
                       JOIN activity_sections x ON x.section_id=e.section_id AND x.activity_id=%s
                       WHERE e.class_id=%s""", (activity_id, class_id))
        roster = {row["student_id"]: row["section_id"] for row in cur.fetchall()}
        if len({item.student_id for item in payload.scores}) != len(payload.scores):
            raise HTTPException(status_code=422, detail="Each student may appear once.")
        for item in payload.scores:
            if item.student_id not in roster or (item.score is not None and item.score > activity["points"]):
                raise HTTPException(status_code=422, detail="Invalid student or score.")
            if item.score is None:
                continue
            cur.execute("""INSERT INTO act_submission(activity_id,student_id,score,submission_date,
                           submission_status,grading_period,graded_at)
                           SELECT %s,%s,%s,CURRENT_TIMESTAMP,'Graded',%s,CURRENT_TIMESTAMP
                           WHERE NOT EXISTS (SELECT 1 FROM act_submission
                             WHERE activity_id=%s AND student_id=%s)""",
                        (activity_id, item.student_id, item.score, activity["grading_period"],
                         activity_id, item.student_id))
            cur.execute("""UPDATE act_submission SET score=%s,submission_status='Graded',
                           graded_at=CURRENT_TIMESTAMP WHERE act_submission_id=(
                             SELECT act_submission_id FROM act_submission
                             WHERE activity_id=%s AND student_id=%s
                             ORDER BY act_submission_id DESC LIMIT 1)""",
                        (item.score, activity_id, item.student_id))
        cur.execute("""DELETE FROM grade_publication WHERE class_id=%s AND section_id IN
                       (SELECT section_id FROM activity_sections WHERE activity_id=%s)""",
                    (class_id, activity_id))
        conn.commit()
        return {"saved": len(payload.scores)}
    finally:
        conn.close()

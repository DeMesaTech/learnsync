"""Versioned faculty syllabus editing, review, and export."""
import io
import json
import uuid
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from psycopg2.extras import Json, RealDictCursor
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.utils import simpleSplit
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from xml.sax.saxutils import escape

from db import get_db_connection
from syllabus_document import import_syllabus
from academic_grading import load_class_preview

syllabus_router = APIRouter(prefix="/api/faculty", tags=["faculty syllabus"])


class DraftInput(BaseModel):
    teacher_id: int
    content: dict
    grading_rule: dict
    acknowledge_warnings: bool = False


class ItemLinkInput(BaseModel):
    teacher_id: int
    item_type: str
    item_id: int
    topic_key: str | None = None


class OverrideInput(BaseModel):
    teacher_id: int
    section_id: int
    item_type: str
    item_id: int
    visible: bool = True
    due_at: datetime | None = None


class PublishInput(BaseModel):
    teacher_id: int
    section_id: int
    grading_period: str


def _owns(cur, class_id: int, teacher_id: int):
    cur.execute("SELECT 1 FROM class WHERE class_id=%s AND employee_id=%s AND workflow_status='active'",
                (class_id, teacher_id))
    if not cur.fetchone():
        raise HTTPException(status_code=403, detail="Assigned faculty member required.")


def _rules(rule):
    weights = rule.get("weights") or {}
    shares = rule.get("period_shares") or {}
    for expected, values in ((("attendance", "quiz", "activity", "exam"), weights),
                             (("Midterm", "Finals"), shares)):
        if set(values) != set(expected) or any(not isinstance(values[key], (int, float)) or
                                                 isinstance(values[key], bool) or values[key] < 0
                                                 for key in expected):
            raise HTTPException(status_code=422, detail="Grading weights and period shares must be nonnegative numbers.")
        if abs(sum(values.values()) - 100) > 0.01:
            raise HTTPException(status_code=422, detail="Each grading weight group must total 100%.")
    if rule.get("transmutation") not in {"raw", "sample_50_100"}:
        raise HTTPException(status_code=422, detail="Choose a supported transmutation rule.")
    threshold = rule.get("passing_threshold")
    if not isinstance(threshold, (int, float)) or isinstance(threshold, bool) or not 0 <= threshold <= 100:
        raise HTTPException(status_code=422, detail="Passing threshold must be between 0 and 100.")
    late = rule.get("attendance_late_percent", 50)
    if not isinstance(late, (int, float)) or isinstance(late, bool) or not 0 <= late <= 100:
        raise HTTPException(status_code=422, detail="Late attendance credit must be between 0 and 100%.")


def _content(content):
    if not str(content.get("title") or "").strip():
        raise HTTPException(status_code=422, detail="Course title is required.")
    chapters = content.get("chapters")
    if not isinstance(chapters, list) or not chapters:
        raise HTTPException(status_code=422, detail="Add at least one syllabus chapter.")
    keys = set()
    def check_node(node, label):
        if not isinstance(node, dict) or not str(node.get("title") or "").strip():
            raise HTTPException(status_code=422, detail=f"Every {label} needs a title.")
        node.setdefault("key", uuid.uuid4().hex)
        if node["key"] in keys:
            raise HTTPException(status_code=422, detail="Chapter, subsection, and topic keys must be unique.")
        keys.add(node["key"])
    for chapter in chapters:
        check_node(chapter, "chapter")
        for topic in chapter.get("topics") or []:
            check_node(topic, "topic")
        for subsection in chapter.get("subsections") or []:
            check_node(subsection, "subsection")
            for topic in subsection.get("topics") or []:
                check_node(topic, "topic")
    return keys


def _version(cur, class_id, status):
    cur.execute(
        "SELECT * FROM syllabus_version WHERE class_id=%s AND status=%s ORDER BY version DESC LIMIT 1",
        (class_id, status),
    )
    return cur.fetchone()


@syllabus_router.get("/{class_id}/syllabus")
def get_syllabus(class_id: int, teacher_id: int):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _owns(cur, class_id, teacher_id)
        draft = _version(cur, class_id, "draft")
        approved = _version(cur, class_id, "approved")
        cur.execute("SELECT file_name,file_path FROM class_syllabus WHERE class_id=%s", (class_id,))
        legacy = cur.fetchone()
        return {"draft": draft, "approved": approved, "legacy_source": legacy}
    finally:
        conn.close()


def _new_draft(cur, class_id, source_name, source_path, content, rule, warnings):
    cur.execute("DELETE FROM syllabus_version WHERE class_id=%s AND status='draft'", (class_id,))
    cur.execute("SELECT COALESCE(MAX(version),0)+1 AS next_version FROM syllabus_version WHERE class_id=%s", (class_id,))
    version = cur.fetchone()["next_version"]
    cur.execute(
        """INSERT INTO syllabus_version(class_id,version,source_name,source_path,content,grading_rule,import_warnings)
           VALUES (%s,%s,%s,%s,%s,%s,%s) RETURNING syllabus_id,version,status""",
        (class_id, version, source_name, source_path, Json(content), Json(rule), Json(warnings)),
    )
    return cur.fetchone()


def _reuse_outline_keys(imported, approved):
    """Keep stable keys when a reimport has the same chapter/topic placement."""
    if not approved:
        return
    previous = {}
    for chapter in approved.get("chapters", []):
        previous.setdefault(str(chapter.get("title", "")).strip().casefold(), []).append(chapter)
    for chapter in imported.get("chapters", []):
        matches = previous.get(str(chapter.get("title", "")).strip().casefold(), [])
        old_chapter = matches.pop(0) if matches else None
        if not old_chapter:
            continue
        chapter["key"] = old_chapter.get("key", chapter.get("key"))
        old_topics = {}
        for topic in old_chapter.get("topics", []):
            old_topics.setdefault(str(topic.get("title", "")).strip().casefold(), []).append(topic)
        for topic in chapter.get("topics", []):
            topic_matches = old_topics.get(str(topic.get("title", "")).strip().casefold(), [])
            old = topic_matches.pop(0) if topic_matches else None
            if old:
                topic["key"] = old.get("key", topic.get("key"))


@syllabus_router.post("/{class_id}/syllabus/import")
async def import_file(class_id: int, teacher_id: int = Form(...), file: UploadFile = File(...)):
    extension = Path(file.filename or "").suffix.lower()
    if extension not in {".docx", ".pdf"}:
        raise HTTPException(status_code=415, detail="Upload a DOCX or text-based PDF.")
    data = await file.read(20 * 1024 * 1024 + 1)
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Syllabus must be 20 MB or smaller.")
    conn = get_db_connection()
    path = None
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _owns(cur, class_id, teacher_id)
        directory = Path(__file__).resolve().parents[1] / "uploads" / "syllabus"
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"{uuid.uuid4().hex}{extension}"
        path.write_bytes(data)
        content, rule, warnings = import_syllabus(path)
        approved = _version(cur, class_id, "approved")
        _reuse_outline_keys(content, approved["content"] if approved else None)
        row = _new_draft(cur, class_id, Path(file.filename).name[:255],
                         f"uploads/syllabus/{path.name}", content, rule, warnings)
        conn.commit()
        return {**row, "content": content, "grading_rule": rule, "import_warnings": warnings}
    except Exception:
        conn.rollback()
        if path:
            path.unlink(missing_ok=True)
        raise
    finally:
        conn.close()


@syllabus_router.post("/{class_id}/syllabus/draft")
def create_draft(class_id: int, teacher_id: int):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _owns(cur, class_id, teacher_id)
        approved = _version(cur, class_id, "approved")
        content = approved["content"] if approved else {"course_code": "", "title": "", "description": "",
                                                   "units": "", "outcomes": [], "chapters": [], "references": []}
        rule = approved["grading_rule"] if approved else {"weights": {"attendance": 0, "quiz": 0,
                    "activity": 0, "exam": 0}, "period_shares": {"Midterm": 50, "Finals": 50},
                    "transmutation": "raw", "passing_threshold": 75,
                    "attendance_late_percent": 50}
        row = _new_draft(cur, class_id, approved["source_name"] if approved else None,
                         approved["source_path"] if approved else None, content, rule, [])
        conn.commit()
        return row
    finally:
        conn.close()


@syllabus_router.put("/{class_id}/syllabus/draft")
def save_draft(class_id: int, payload: DraftInput):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _owns(cur, class_id, payload.teacher_id)
        row = _version(cur, class_id, "draft")
        if not row:
            raise HTTPException(status_code=404, detail="Create or import a draft first.")
        cur.execute(
            """UPDATE syllabus_version SET content=%s,grading_rule=%s,warnings_acknowledged=%s
               WHERE syllabus_id=%s""",
            (Json(payload.content), Json(payload.grading_rule), payload.acknowledge_warnings, row["syllabus_id"]),
        )
        conn.commit()
        return {"saved": True, "version": row["version"]}
    finally:
        conn.close()


@syllabus_router.post("/{class_id}/syllabus/approve")
def approve(class_id: int, teacher_id: int):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _owns(cur, class_id, teacher_id)
        draft = _version(cur, class_id, "draft")
        if not draft:
            raise HTTPException(status_code=404, detail="No draft to approve.")
        valid_keys = _content(draft["content"])
        _rules(draft["grading_rule"])
        cur.execute("UPDATE syllabus_version SET content=%s WHERE syllabus_id=%s",
                    (Json(draft["content"]), draft["syllabus_id"]))
        if draft["import_warnings"] and not draft["warnings_acknowledged"]:
            raise HTTPException(status_code=422, detail="Review and acknowledge import warnings before approval.")
        cur.execute("SELECT topic_key FROM syllabus_item_link WHERE class_id=%s AND topic_key IS NOT NULL", (class_id,))
        linked = {row["topic_key"] for row in cur.fetchall()}
        cur.execute("SELECT DISTINCT topic_key FROM learning_content WHERE class_id=%s", (class_id,))
        linked.update(row["topic_key"] for row in cur.fetchall())
        cur.execute("SELECT DISTINCT content_key FROM quiz WHERE class_id=%s AND content_key IS NOT NULL", (class_id,))
        linked.update(row["content_key"] for row in cur.fetchall())
        cur.execute("SELECT DISTINCT content_key FROM activity WHERE class_id=%s AND content_key IS NOT NULL", (class_id,))
        linked.update(row["content_key"] for row in cur.fetchall())
        cur.execute("SELECT DISTINCT content_key FROM learning_resource WHERE class_id=%s", (class_id,))
        linked.update(row["content_key"] for row in cur.fetchall())
        if linked - valid_keys:
            raise HTTPException(status_code=422, detail="Reassign linked lessons, materials, activities, and quizzes before approving the revised outline.")
        cur.execute("UPDATE syllabus_version SET status='superseded' WHERE class_id=%s AND status='approved'", (class_id,))
        cur.execute("UPDATE syllabus_version SET status='approved',approved_at=CURRENT_TIMESTAMP WHERE syllabus_id=%s",
                    (draft["syllabus_id"],))
        cur.execute("DELETE FROM grade_publication WHERE class_id=%s", (class_id,))
        cur.execute("UPDATE grade_visibility SET visible=false WHERE class_id=%s", (class_id,))
        # Keep existing chapter IDs where titles match so module/progress links survive a rule revision.
        for order, chapter in enumerate(draft["content"]["chapters"], start=1):
            cur.execute("SELECT topic_id FROM syllabus_topic WHERE class_id=%s AND lower(title)=lower(%s) LIMIT 1",
                        (class_id, chapter["title"]))
            existing = cur.fetchone()
            if existing:
                cur.execute("UPDATE syllabus_topic SET display_order=%s WHERE topic_id=%s",
                            (order, existing["topic_id"]))
            else:
                cur.execute("INSERT INTO syllabus_topic(class_id,title,display_order) VALUES (%s,%s,%s)",
                            (class_id, chapter["title"], order))
        conn.commit()
        return {"approved": True, "version": draft["version"], "syllabus_id": draft["syllabus_id"]}
    finally:
        conn.close()


def _document_lines(content, rule):
    yield "COURSE SYLLABUS"
    yield f"{content.get('course_code', '')}  {content.get('title', '')}".strip()
    yield f"Credit units: {content.get('units', '')}"
    yield f"Description: {content.get('description', '')}"
    yield "Course learning outcomes"
    for outcome in content.get("outcomes") or []:
        yield f"• {outcome}"
    yield "Course coverage"
    for chapter in content.get("chapters") or []:
        yield f"{chapter.get('week', '')}  {chapter.get('title', '')}".strip()
        for topic in chapter.get("topics") or []:
            yield f"  • {topic.get('title', '')}"
            for label, field in (("Outcomes", "outcomes"), ("Learning materials", "materials"),
                                 ("References", "references")):
                if topic.get(field):
                    yield f"    {label}: {', '.join(topic[field])}"
        for subsection in chapter.get("subsections") or []:
            yield f"  Subsection: {subsection.get('title', '')}"
            for topic in subsection.get("topics") or []:
                yield f"    • {topic.get('title', '')}"
                for label, field in (("Outcomes", "outcomes"), ("Learning materials", "materials"),
                                     ("References", "references")):
                    if topic.get(field):
                        yield f"      {label}: {', '.join(topic[field])}"
        if chapter.get("assessments"):
            yield f"Assessments: {', '.join(chapter['assessments'])}"
    yield "References"
    for reference in content.get("references") or []:
        yield f"• {reference}"
    yield "Grading rules"
    yield "Category weights: " + ", ".join(f"{key} {value}%" for key, value in rule["weights"].items())
    yield "Period shares: " + ", ".join(f"{key} {value}%" for key, value in rule["period_shares"].items())
    yield f"Transmutation: {rule['transmutation']}; Passing threshold: {rule['passing_threshold']}"
    yield f"Late attendance credit: {rule.get('attendance_late_percent', 50)}%; Excused dates excluded"


@syllabus_router.get("/{class_id}/syllabus/export/{format}")
def export_syllabus(class_id: int, format: str, teacher_id: int):
    if format not in {"docx", "pdf"}:
        raise HTTPException(status_code=404, detail="Supported exports: docx, pdf.")
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _owns(cur, class_id, teacher_id)
        approved = _version(cur, class_id, "approved")
        if not approved:
            raise HTTPException(status_code=409, detail="Approve the syllabus before exporting.")
        lines = list(_document_lines(approved["content"], approved["grading_rule"]))
    finally:
        conn.close()
    stream = io.BytesIO()
    if format == "docx":
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.8)
        document.add_heading("Governor Mariano E. Villafuerte Community College", 0)
        document.add_paragraph("Approved course syllabus", style="Subtitle")
        for index, line in enumerate(lines):
            if index in {0, 4} or line in {"Course coverage", "References", "Grading rules"}:
                document.add_heading(line, level=1)
            else:
                document.add_paragraph(line)
        document.save(stream)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        pdf = SimpleDocTemplate(stream, pagesize=(595, 842), rightMargin=48, leftMargin=48,
                                topMargin=48, bottomMargin=48)
        styles = getSampleStyleSheet()
        story = [Paragraph("Governor Mariano E. Villafuerte Community College", styles["Title"]), Spacer(1, 16)]
        for index, line in enumerate(lines):
            style = styles["Heading2"] if index in {0, 4} or line in {"Course coverage", "References", "Grading rules"} else styles["BodyText"]
            story.extend([Paragraph(escape(line), style), Spacer(1, 7)])
        pdf.build(story)
        mime = "application/pdf"
    stream.seek(0)
    return StreamingResponse(stream, media_type=mime,
        headers={"Content-Disposition": f'attachment; filename="learnsync-syllabus-{class_id}-v{approved["version"]}.{format}"'})


@syllabus_router.put("/{class_id}/syllabus/item-link")
def link_item(class_id: int, payload: ItemLinkInput):
    if payload.item_type not in {"module", "activity", "quiz"}:
        raise HTTPException(status_code=422, detail="Invalid item type.")
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _owns(cur, class_id, payload.teacher_id)
        cur.execute(f"SELECT 1 FROM {payload.item_type} WHERE {payload.item_type}_id=%s AND class_id=%s",
                    (payload.item_id, class_id))
        if not cur.fetchone():
            raise HTTPException(status_code=404, detail="Content item not found.")
        if payload.topic_key:
            approved = _version(cur, class_id, "approved")
            keys = {topic.get("key") for chapter in (approved["content"].get("chapters") if approved else [])
                    for topic in chapter.get("topics") or []}
            keys.update(topic.get("key") for chapter in (approved["content"].get("chapters") if approved else [])
                        for subsection in chapter.get("subsections") or []
                        for topic in subsection.get("topics") or [])
            if payload.topic_key not in keys:
                raise HTTPException(status_code=422, detail="Topic not found in the approved syllabus.")
            cur.execute("""INSERT INTO syllabus_item_link(class_id,item_type,item_id,topic_key)
                           VALUES (%s,%s,%s,%s) ON CONFLICT(class_id,item_type,item_id)
                           DO UPDATE SET topic_key=EXCLUDED.topic_key""",
                        (class_id, payload.item_type, payload.item_id, payload.topic_key))
        else:
            cur.execute("DELETE FROM syllabus_item_link WHERE class_id=%s AND item_type=%s AND item_id=%s",
                        (class_id, payload.item_type, payload.item_id))
        conn.commit()
        return {"linked": bool(payload.topic_key)}
    finally:
        conn.close()


@syllabus_router.get("/{class_id}/syllabus/items")
def syllabus_items(class_id: int, teacher_id: int):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _owns(cur, class_id, teacher_id)
        cur.execute("""SELECT s.section_id,s.section FROM class_section cs
                       JOIN section s ON s.section_id=cs.section_id WHERE cs.class_id=%s ORDER BY s.section""",
                    (class_id,))
        sections = cur.fetchall()
        items = []
        for table, id_column in (("module", "module_id"), ("activity", "activity_id"), ("quiz", "quiz_id")):
            cur.execute(f"SELECT {id_column} AS item_id,title FROM {table} WHERE class_id=%s ORDER BY {id_column}",
                        (class_id,))
            items.extend({"item_type": table, **row} for row in cur.fetchall())
        cur.execute("SELECT item_type,item_id,topic_key FROM syllabus_item_link WHERE class_id=%s", (class_id,))
        links = cur.fetchall()
        cur.execute("""SELECT section_id,item_type,item_id,visible,due_at FROM syllabus_section_override
                       WHERE class_id=%s""", (class_id,))
        overrides = cur.fetchall()
        return {"sections": sections, "items": items, "links": links, "overrides": overrides}
    finally:
        conn.close()


@syllabus_router.put("/{class_id}/syllabus/section-override")
def set_override(class_id: int, payload: OverrideInput):
    if payload.item_type not in {"module", "activity", "quiz"}:
        raise HTTPException(status_code=422, detail="Invalid item type.")
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _owns(cur, class_id, payload.teacher_id)
        cur.execute("SELECT 1 FROM class_section WHERE class_id=%s AND section_id=%s",
                    (class_id, payload.section_id))
        if not cur.fetchone():
            raise HTTPException(status_code=404, detail="Section is not assigned to this subject.")
        cur.execute(f"SELECT 1 FROM {payload.item_type} WHERE {payload.item_type}_id=%s AND class_id=%s",
                    (payload.item_id, class_id))
        if not cur.fetchone():
            raise HTTPException(status_code=404, detail="Content item not found.")
        cur.execute("""INSERT INTO syllabus_section_override(class_id,section_id,item_type,item_id,visible,due_at)
                       VALUES (%s,%s,%s,%s,%s,%s) ON CONFLICT(class_id,section_id,item_type,item_id)
                       DO UPDATE SET visible=EXCLUDED.visible,due_at=EXCLUDED.due_at""",
                    (class_id, payload.section_id, payload.item_type, payload.item_id,
                     payload.visible, payload.due_at))
        conn.commit()
        return {"saved": True}
    finally:
        conn.close()


@syllabus_router.get("/{class_id}/grades/preview")
def grade_preview(class_id: int, teacher_id: int, section_id: int):
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _owns(cur, class_id, teacher_id)
        cur.execute("SELECT 1 FROM class_section WHERE class_id=%s AND section_id=%s", (class_id, section_id))
        if not cur.fetchone():
            raise HTTPException(status_code=404, detail="Section not assigned to this subject.")
        approved = _version(cur, class_id, "approved")
        if not approved:
            raise HTTPException(status_code=409, detail="Approve a syllabus before calculating grades.")
        return load_class_preview(cur, class_id, section_id, approved)
    finally:
        conn.close()


@syllabus_router.post("/{class_id}/grades/publish")
def publish_grades(class_id: int, payload: PublishInput):
    if payload.grading_period not in {"Midterm", "Finals", "Course"}:
        raise HTTPException(status_code=422, detail="Invalid grading period.")
    conn = get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        _owns(cur, class_id, payload.teacher_id)
        cur.execute("SELECT 1 FROM class_section WHERE class_id=%s AND section_id=%s", (class_id, payload.section_id))
        if not cur.fetchone():
            raise HTTPException(status_code=404, detail="Section not assigned to this subject.")
        approved = _version(cur, class_id, "approved")
        if not approved:
            raise HTTPException(status_code=409, detail="Approve a syllabus first.")
        preview = load_class_preview(cur, class_id, payload.section_id, approved)
        if not preview["students"]:
            raise HTTPException(status_code=409, detail="Section has no students.")
        for student in preview["students"]:
            grade = (student["course_grade"] if payload.grading_period == "Course" else
                     student["periods"][payload.grading_period]["grade"])
            if grade is None:
                raise HTTPException(status_code=409, detail="All required scores must be complete before publication.")
        cur.execute(
            """INSERT INTO grade_publication(class_id,section_id,grading_period,syllabus_id,snapshot,published_by)
               VALUES (%s,%s,%s,%s,%s,%s)
               ON CONFLICT(class_id,section_id,grading_period) DO UPDATE
               SET syllabus_id=EXCLUDED.syllabus_id,snapshot=EXCLUDED.snapshot,
                   published_by=EXCLUDED.published_by,published_at=CURRENT_TIMESTAMP""",
            (class_id, payload.section_id, payload.grading_period, approved["syllabus_id"],
             Json(preview), payload.teacher_id),
        )
        conn.commit()
        return {"published": True, "syllabus_version": approved["version"]}
    finally:
        conn.close()

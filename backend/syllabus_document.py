"""Import course syllabus content into the faculty review model."""
import re
import uuid
from pathlib import Path

from docx import Document
from pypdf import PdfReader


CHAPTER = re.compile(r"\bCHAPTER\s+(\d+)\s*[-–—:]?\s*([^\n]*)", re.I)
PERCENT = re.compile(r"(attendance|participation|recitation|quizzes|projects?|activities|examination|exam)\s*[-:. ]+\s*(\d+(?:\.\d+)?)\s*%", re.I)


def _clean(value):
    return re.sub(r"\s+", " ", value or "").strip(" -–—\t")


def _items(value):
    return [_clean(line) for line in re.split(r"\n+", value or "") if _clean(line)]


def _chapter_from_text(text):
    match = CHAPTER.search(text or "")
    if not match:
        return None
    lines = _items(text)
    title = _clean(f"Chapter {match.group(1)} {match.group(2)}")
    match_index = next((i for i, line in enumerate(lines) if CHAPTER.search(line)), 0)
    following = [line for line in lines[match_index + 1:] if not CHAPTER.search(line)]
    return {"key": uuid.uuid4().hex, "title": title, "week": "", "topics": [
        {"key": uuid.uuid4().hex, "title": line, "outcomes": [], "materials": [], "references": []}
        for line in following if len(line) < 200
    ], "outcomes": [], "materials": [], "assessments": [], "references": []}


def _rules(text):
    matches = {name.lower(): float(value) for name, value in PERCENT.findall(text)}
    practical = re.search(r"Projects?/Practical Exercises[^\n%]*?(\d+(?:\.\d+)?)\s*%", text, re.I)
    participation = matches.get("participation", 0) + matches.get("recitation", 0)
    activity = participation + matches.get("project", 0) + matches.get("projects", 0) + matches.get("activities", 0)
    if practical and not matches.get("project") and not matches.get("projects"):
        activity += float(practical.group(1))
    quiz = matches.get("quizzes", 0)
    exam = matches.get("examination", 0) + matches.get("exam", 0)
    weights = {"attendance": matches.get("attendance", 0), "quiz": quiz,
               "activity": activity, "exam": exam}
    return {"weights": weights, "period_shares": {"Midterm": 50, "Finals": 50},
            "transmutation": "raw", "passing_threshold": 75,
            "attendance_late_percent": 50}


def _parse_docx(path):
    doc = Document(path)
    content = {"course_code": "", "title": "", "description": "", "units": "",
               "outcomes": [], "chapters": [], "references": []}
    warnings = []
    all_text = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            all_text.append(paragraph.text)
    coverage = False
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            unique = list(dict.fromkeys(cells))
            all_text.extend(unique)
            label = _clean(unique[0] if unique else "").lower()
            value = _clean(unique[-1] if unique else "")
            if "course no" in label:
                content["course_code"] = value
            elif "course name" in label:
                content["title"] = value
            elif "course description" in label:
                content["description"] = value
            elif "credit units" in label:
                content["units"] = value
            if "course learning outcomes" in label:
                coverage = False
            elif "course coverage" in label:
                coverage = True
            elif label.startswith("11. references") or label.startswith("references"):
                content["references"].extend(_items(unique[0])[1:])
            elif coverage and re.search(r"\bweek\s*\d+", label, re.I):
                topic_cell = next((c for c in unique if CHAPTER.search(c)), "")
                chapter = _chapter_from_text(topic_cell)
                if chapter:
                    chapter["week"] = _clean(unique[0])
                    chapter["outcomes"] = _items(unique[1]) if len(unique) > 1 else []
                    chapter["assessments"] = _items(unique[-1]) if len(unique) > 2 else []
                    chapter["materials"] = _items(unique[-2]) if len(unique) > 3 else []
                    content["chapters"].append(chapter)
            elif len(unique) > 1 and unique[0].strip().startswith("•"):
                content["outcomes"].append(_clean(unique[0]))
    if not content["outcomes"]:
        warnings.append("Course outcomes need review; the document's outcome matrix could not be mapped confidently.")
    return content, _rules("\n".join(all_text)), warnings


def _parse_pdf(path):
    reader = PdfReader(path)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if not text.strip():
        return {"course_code": "", "title": "", "description": "", "units": "",
                "outcomes": [], "chapters": [], "references": []}, _rules(""), [
                "This PDF has no selectable text. Complete the syllabus manually or import a DOCX."]
    content = {"course_code": "", "title": "", "description": "", "units": "",
               "outcomes": [], "chapters": [], "references": []}
    for field, pattern in (("course_code", r"Course\s+No\.?\s*[:\-]?\s*(.+)"),
                           ("title", r"Course\s+Name\s*[:\-]?\s*(.+)"),
                           ("description", r"Course\s+Description\s*[:\-]?\s*(.+)"),
                           ("units", r"Credit\s+Units\s*[:\-]?\s*(.+)")):
        match = re.search(pattern, text, re.I)
        if match:
            content[field] = _clean(match.group(1))[:500]
    seen = set()
    for line in text.splitlines():
        chapter = _chapter_from_text(line)
        if chapter and chapter["title"].casefold() not in seen:
            seen.add(chapter["title"].casefold())
            content["chapters"].append(chapter)
    return content, _rules(text), ["Review outcomes, topics, materials, and references extracted from this PDF."]


def import_syllabus(path: str | Path):
    path = Path(path)
    if path.suffix.lower() == ".docx":
        content, rules, warnings = _parse_docx(path)
    elif path.suffix.lower() == ".pdf":
        content, rules, warnings = _parse_pdf(path)
    else:
        raise ValueError("Only DOCX and text-based PDF are supported.")
    if not content["title"]:
        warnings.append("Course title is missing.")
    if not content["chapters"]:
        warnings.append("No chapters were detected; add them before approval.")
    if abs(sum(rules["weights"].values()) - 100) > 0.01:
        warnings.append("Assessment weights need faculty review and must total 100%.")
    return content, rules, warnings
